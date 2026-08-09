"""Границы, за которыми внешние данные ломали обработчик.

Всё здесь — про значения, приходящие снаружи: callback_data подделывается
клиентом, аргумент /allow набирается руками, имя ученика придумывает человек.
Общий признак дефектов был один: проверялась форма, но не величина и не
содержимое, — и падение случалось уже в SQLite или в Excel у получателя.
"""

import io

import openpyxl
from sqlalchemy import select

from lesson_tracker.bot.export_data import rows_to_xlsx
from lesson_tracker.bot.handlers._common import as_int
from lesson_tracker.db.models import AllowedUser, Student
from lesson_tracker.services import PaymentService, StudentService
from tests.bot_harness import make_update_callback

ADMIN = 1
HUGE = "9" * 23  # влезает в лимит callback_data, но не в 64-битный INTEGER


def test_as_int_отвергает_число_шире_базы():
    assert as_int("5") == 5
    assert as_int("не число") is None
    assert as_int(HUGE) is None, "SQLite такое не примет — отвергать надо до него"


async def test_подделанный_огромный_id_не_роняет_обработчик(harness, session, sessionmaker):
    """Раньше 23-значный id проходил разбор и падал на привязке параметра:
    пользователь видел «Не получилось выполнить действие» и трейсбек в
    журнале вместо «Кнопка устарела»."""
    await StudentService(session).create(ADMIN, "Аня", 160000)
    await session.commit()

    await harness.dp.feed_update(
        harness.bot, make_update_callback(f"card:{HUGE}", user_id=ADMIN, update_id=9500)
    )

    answers = [m.text or "" for m in harness.session.calls_of("AnswerCallbackQuery")]
    assert any("устарела" in t for t in answers), f"ожидался внятный отказ, а было: {answers}"


async def test_огромный_id_в_allow_даёт_подсказку(harness, sessionmaker):
    await harness.send(f"/allow {HUGE}", user_id=ADMIN)

    sent = [m.text or "" for m in harness.session.calls_of("SendMessage")]
    assert any("Использование" in t for t in sent), f"ожидалась подсказка, а было: {sent}"
    async with sessionmaker() as s:
        assert list(await s.scalars(select(AllowedUser))) == []


def test_выгрузка_в_excel_не_исполняет_формулы():
    """Имя, начинающееся с «=», openpyxl записывает настоящей формулой:
    выгрузка превращалась в исполняемый документ, а кривое выражение — в
    предложение Excel «восстановить файл». Экран стоял только на пути CSV."""
    data = rows_to_xlsx([['=HYPERLINK("http://evil","клик")', "ок"]], "Лист")
    cell = openpyxl.load_workbook(io.BytesIO(data)).active.cell(row=1, column=1)

    assert cell.data_type == "s", "ячейка не должна быть формулой"


async def test_нулевая_стоимость_не_роняет_оплату(session):
    """Инвариант «стоимость больше нуля» держался только в обработчиках, а
    делит на неё сервис. Боевая база приехала из serverless-версии, так что
    унаследованная строка с нулём — не выдумка."""
    students = StudentService(session)
    assert await students.create(ADMIN, "Аня", 0) is None, "ноль нельзя принимать"

    session.add(Student(owner_id=ADMIN, name="Боря", name_lower="боря", price=0))
    await session.flush()
    broken = (await session.scalars(select(Student).where(Student.name == "Боря"))).one()

    assert await PaymentService(session, students).apply(ADMIN, broken.id, 1000) is None


async def test_гигантский_поисковый_запрос_не_пробивает_лимит_telegram(harness, session):
    """Запрос — пользовательский текст до 4096 символов. Эхо без обрезки
    давало сообщение длиннее лимита: API отвечал 400, ответ молча терялся, и
    бот выглядел мёртвым — подсказка удалена, состояние очищено, в чате пусто.
    """
    await StudentService(session).create(ADMIN, "Аня", 160000)
    await session.commit()

    await harness.click("search", user_id=ADMIN)
    await harness.send("х" * 4096, user_id=ADMIN)

    sent = [m.text or "" for m in harness.session.calls_of("SendMessage")]
    result = [t for t in sent if "🔍" in t]
    assert result, "ответ на поиск обязан уйти"
    assert all(len(t) <= 4096 for t in result), f"эхо пробило лимит: {max(map(len, result))}"


def test_карточка_переживает_null_поля_последней_оплаты():
    """База пришла живой из serverless (L1): строка с датой оплаты, но
    NULL-суммой — не гипотеза. format_money(None) ронял карточку навсегда."""
    from lesson_tracker.bot.render import render_card

    class S:
        name = "Аня"
        balance = 3
        price = 160000
        remainder = 0
        last_payment_at = 1_700_000_000
        last_payment_amount = None
        last_payment_lessons = None

    card = render_card(S())
    assert "Последняя оплата" in card
    assert "None" not in card


def test_история_переживает_снимок_без_цены():
    """snapshot_before — нетипизированный JSON из serverless; отсутствие ключа
    price роняло всю «Историю» ученика навсегда."""
    from lesson_tracker.bot.render import render_operation

    class Op:
        type = "price_change"
        lessons_delta = 0
        balance_after = 0
        remainder_after = 0
        amount = None
        new_price = 180000
        created_at = 1_700_000_000
        undone = False
        snapshot_before: dict = {}

    line = render_operation(Op())
    assert "?" in line and "1 800 ₽" in line


def test_числовые_ячейки_xlsx_остаются_числами():
    """guard_formula строковал всё подряд: =SUM() по выгрузке возвращал 0,
    сортировка по «Осталось занятий» ставила 10 перед 4, а Excel зеленил
    таблицу флагом «число как текст». Число формулой не станет — его не
    трогаем; формульные строки по-прежнему экранируются."""
    data = rows_to_xlsx([[7, "Аня", 10], ["=1+1", "ок", -3]], "Лист")
    ws = openpyxl.load_workbook(io.BytesIO(data)).active

    assert ws.cell(row=1, column=1).data_type == "n", "int обязан остаться числом"
    assert ws.cell(row=1, column=3).data_type == "n"
    assert ws.cell(row=2, column=3).data_type == "n", "отрицательное число — тоже число"
    assert ws.cell(row=2, column=1).data_type == "s", "формула обязана остаться текстом"
